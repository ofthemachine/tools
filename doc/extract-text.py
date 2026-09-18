#!/usr/bin/env -S fragletc --image ofthemachine/python3@sha256:a015b4f9f0e7c1648f8fcd4164f6a296dd3c9189404f88f333a9764d08d13522
#: d=Extract readable text out of PDF, DOCX, PPTX, XLSX, ODT, RTF, HTML, and plain-text documents into structured Markdown carrying page, slide, and sheet markers. The format is detected from file content, so no extension is needed; pages selects a 1-indexed subset such as "1-5" or "2,4,7-9" (PDF pages, PPTX slides, XLSX sheets).
#: when=Use when the user wants the text of a PDF, Word, PowerPoint, Excel, ODT, RTF, or HTML file, or a document's page count.
#: network=none
#: stdin=none
#: category=extraction
#: param=document:required:file:d=Document mounted at /input/document; format sniffed from content
#: param=pages:d=1-indexed subset like 1-5 or 2,4,7-9 (pages, slides, or sheets)
#: param=format:d=Force a format instead of sniffing: pdf, docx, pptx, xlsx, odt, rtf, html, text
import os
import sys
import zipfile

path = os.environ.get("DOCUMENT", "")
pages_spec = os.environ.get("PAGES", "").strip()
forced_format = os.environ.get("FORMAT", "").strip().lower()

if not path or not os.path.exists(path):
    print(f"Error: document file '{path}' does not exist", file=sys.stderr)
    sys.exit(1)

with open(path, "rb") as f:
    head = f.read(8)


def detect_format() -> str:
    if head.startswith(b"%PDF"):
        return "pdf"
    if head.startswith(b"{\\rtf"):
        return "rtf"
    if head.startswith(b"PK\x03\x04"):
        try:
            names = set(zipfile.ZipFile(path).namelist())
        except zipfile.BadZipFile:
            return "text"
        if "word/document.xml" in names:
            return "docx"
        if "ppt/presentation.xml" in names:
            return "pptx"
        if "xl/workbook.xml" in names:
            return "xlsx"
        if "content.xml" in names:
            return "odt"
        return "text"
    sample = head.lstrip().lower()
    if sample.startswith(b"<!doc") or sample.startswith(b"<html"):
        return "html"
    return "text"


fmt = forced_format or detect_format()


def parse_pages(spec: str):
    """1-indexed selection like "1-5" or "2,4,7-9"; None means everything."""
    if not spec:
        return None
    wanted = set()
    for part in spec.split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            lo, _, hi = part.partition("-")
            try:
                wanted.update(range(int(lo), int(hi) + 1))
            except ValueError:
                print(f"Error: bad pages range '{part}'", file=sys.stderr)
                sys.exit(1)
        else:
            try:
                wanted.add(int(part))
            except ValueError:
                print(f"Error: bad pages value '{part}'", file=sys.stderr)
                sys.exit(1)
    return wanted or None


selection = parse_pages(pages_spec)


def selected(index_1based: int) -> bool:
    return selection is None or index_1based in selection


def warn_pages_ignored():
    if selection is not None:
        print(f"pages is not meaningful for {fmt} input; extracting the whole document", file=sys.stderr)


def emit_pdf():
    import pdfplumber

    with pdfplumber.open(path) as pdf:
        total = len(pdf.pages)
        for i, page in enumerate(pdf.pages, start=1):
            if not selected(i):
                continue
            print(f"## Page {i}/{total}\n")
            print((page.extract_text() or "").strip() + "\n")


def emit_docx():
    import docx

    warn_pages_ignored()
    document = docx.Document(path)
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower()
        if style.startswith("heading"):
            level = "".join(c for c in style if c.isdigit()) or "1"
            print(f"{'#' * min(int(level) + 1, 6)} {text}\n")
        else:
            print(text + "\n")
    for t_index, table in enumerate(document.tables, start=1):
        print(f"### Table {t_index}\n")
        for r_index, row in enumerate(table.rows):
            cells = [cell.text.strip() for cell in row.cells]
            print("| " + " | ".join(cells) + " |")
            if r_index == 0:
                print("| " + " | ".join("---" for _ in cells) + " |")
        print()


def emit_pptx():
    from pptx import Presentation

    prs = Presentation(path)
    slides = list(prs.slides)
    for i, slide in enumerate(slides, start=1):
        if not selected(i):
            continue
        print(f"## Slide {i}/{len(slides)}\n")
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    print(text + "\n")
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                print(f"> Speaker notes: {notes}\n")


def emit_xlsx():
    import openpyxl
    from tabulate import tabulate

    # A file handle, not the path: openpyxl validates the extension of a path
    # argument, and the mount at /input/document deliberately has none. The
    # handle must outlive the row iteration, which read_only mode does lazily.
    with open(path, "rb") as handle:
        book = openpyxl.load_workbook(handle, data_only=True, read_only=True)
        sheets = book.sheetnames
        for i, name in enumerate(sheets, start=1):
            if not selected(i):
                continue
            sheet = book[name]
            rows = [["" if c is None else str(c) for c in row] for row in sheet.iter_rows(values_only=True)]
            rows = [r for r in rows if any(cell.strip() for cell in r)]
            print(f"## Sheet {i}/{len(sheets)}: {name}\n")
            if not rows:
                print("(empty)\n")
                continue
            print(tabulate(rows[1:], headers=rows[0], tablefmt="github") + "\n")


def emit_odt():
    from odf import teletype, text as odftext
    from odf.opendocument import load

    warn_pages_ignored()
    doc = load(path)
    for para in doc.getElementsByType(odftext.P):
        line = teletype.extractText(para).strip()
        if line:
            print(line + "\n")


def emit_rtf():
    from striprtf.striprtf import rtf_to_text

    warn_pages_ignored()
    with open(path, "r", errors="ignore") as f:
        print(rtf_to_text(f.read()).strip())


def emit_html():
    import trafilatura

    warn_pages_ignored()
    with open(path, "r", errors="ignore") as f:
        raw = f.read()
    extracted = trafilatura.extract(raw, output_format="markdown")
    if extracted:
        print(extracted.strip())
    else:
        print("Error: no readable article content found in HTML", file=sys.stderr)
        sys.exit(1)


def emit_text():
    warn_pages_ignored()
    with open(path, "r", errors="ignore") as f:
        sys.stdout.write(f.read())


handlers = {
    "pdf": emit_pdf,
    "docx": emit_docx,
    "pptx": emit_pptx,
    "xlsx": emit_xlsx,
    "odt": emit_odt,
    "rtf": emit_rtf,
    "html": emit_html,
    "text": emit_text,
}

handler = handlers.get(fmt)
if handler is None:
    print(f"Error: unsupported format '{fmt}' (expected one of {', '.join(sorted(handlers))})", file=sys.stderr)
    sys.exit(1)

handler()
