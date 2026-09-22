#!/usr/bin/env -S fragletc --image ofthemachine/python3@sha256:a015b4f9f0e7c1648f8fcd4164f6a296dd3c9189404f88f333a9764d08d13522
#: d=Inspect spreadsheet workbooks (XLSX, ODS), listing sheet names, row and column dimensions, and header columns.
#: when=Use when the user wants to see the structure, sheets, dimensions, or column names of an Excel or ODS spreadsheet without reading all rows.
#: network=none
#: stdin=none
#: param=workbook:required:file:d=Spreadsheet workbook mounted at /input/workbook
#: param=format:default=markdown:d=Output format: markdown or json
import json
import os
import sys
import zipfile

import openpyxl
from tabulate import tabulate

wb_path = os.environ.get("WORKBOOK", "")
fmt = os.environ["FORMAT"].strip().lower()

if not wb_path or not os.path.exists(wb_path):
    print(f"Error: workbook file '{wb_path}' does not exist", file=sys.stderr)
    sys.exit(1)

if fmt not in ("markdown", "json"):
    print(f"Error: unknown format '{fmt}' (expected markdown or json)", file=sys.stderr)
    sys.exit(1)


def is_ods(path: str) -> bool:
    try:
        with zipfile.ZipFile(path) as zf:
            return "content.xml" in zf.namelist() and "mimetype" in zf.namelist()
    except Exception:
        return False


sheets_info = []

if is_ods(wb_path):
    from odf import teletype, table
    from odf.opendocument import load

    doc = load(wb_path)
    for idx, sheet in enumerate(doc.getElementsByType(table.Table), start=1):
        name = sheet.getAttribute("name") or f"Sheet{idx}"
        rows = sheet.getElementsByType(table.TableRow)
        header = []
        for r in rows:
            cells = [teletype.extractText(c).strip() for c in r.getElementsByType(table.TableCell)]
            if any(cells):
                header = [c or f"Col_{i+1}" for i, c in enumerate(cells)]
                break
        sheets_info.append({
            "index": idx,
            "name": name,
            "rows": len(rows),
            "columns": len(header),
            "headers": header,
        })
else:
    try:
        with open(wb_path, "rb") as handle:
            wb = openpyxl.load_workbook(handle, data_only=True, read_only=True)
            for idx, name in enumerate(wb.sheetnames, start=1):
                ws = wb[name]
                header = []
                row_count = ws.max_row or 0
                col_count = ws.max_column or 0
                for r in ws.iter_rows(values_only=True):
                    str_cells = ["" if c is None else str(c).strip() for c in r]
                    if any(str_cells):
                        header = [c or f"Col_{i+1}" for i, c in enumerate(str_cells)]
                        break
                sheets_info.append({
                    "index": idx,
                    "name": name,
                    "rows": row_count,
                    "columns": col_count or len(header),
                    "headers": header,
                })
    except Exception as e:
        print(f"Error: failed to open workbook as Excel XLSX or ODS: {e}", file=sys.stderr)
        sys.exit(1)

if not sheets_info:
    print("Workbook has no sheets.", file=sys.stderr)
    sys.exit(1)

if fmt == "json":
    print(json.dumps(sheets_info, indent=2))
else:
    table_rows = []
    for s in sheets_info:
        hdr_preview = ", ".join(s["headers"][:6])
        if len(s["headers"]) > 6:
            hdr_preview += f", ... (+{len(s['headers']) - 6} more)"
        table_rows.append([
            s["index"],
            s["name"],
            s["rows"],
            s["columns"],
            hdr_preview or "(empty)"
        ])
    print(tabulate(table_rows, headers=["#", "Sheet Name", "Rows", "Cols", "Header Preview"], tablefmt="github"))
