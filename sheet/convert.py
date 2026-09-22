#!/usr/bin/env -S fragletc --image ofthemachine/python3@sha256:a015b4f9f0e7c1648f8fcd4164f6a296dd3c9189404f88f333a9764d08d13522
#: d=Convert a sheet from an Excel (XLSX) or ODS spreadsheet into clean CSV, TSV, or JSON.
#: when=Use when the user wants to extract data from an Excel or ODS workbook into a flat format (CSV, TSV, JSON) or save a sheet as a table.
#: network=none
#: stdin=none
#: param=workbook:required:file:d=Spreadsheet workbook mounted at /input/workbook
#: param=sheet:default=1:d=1-indexed sheet index or sheet name to convert (defaults to the first sheet)
#: param=format:default=csv:d=Output format: csv, tsv, json
#: output=sheet.csv
#: output=sheet.tsv
#: output=sheet.json
import csv
import io
import json
import os
import sys
import zipfile

import openpyxl

wb_path = os.environ.get("WORKBOOK", "")
sheet_spec = os.environ["SHEET"].strip()
fmt = os.environ["FORMAT"].strip().lower()

if not wb_path or not os.path.exists(wb_path):
    print(f"Error: workbook file '{wb_path}' does not exist", file=sys.stderr)
    sys.exit(1)

if fmt not in ("csv", "tsv", "json"):
    print(f"Error: unknown format '{fmt}' (expected csv, tsv, json)", file=sys.stderr)
    sys.exit(1)


def is_ods(path: str) -> bool:
    try:
        with zipfile.ZipFile(path) as zf:
            return "content.xml" in zf.namelist() and "mimetype" in zf.namelist()
    except Exception:
        return False


all_rows = []

if is_ods(wb_path):
    from odf import teletype, table
    from odf.opendocument import load

    doc = load(wb_path)
    tables = doc.getElementsByType(table.Table)
    target_table = None

    if sheet_spec.isdigit():
        idx = int(sheet_spec)
        if 1 <= idx <= len(tables):
            target_table = tables[idx - 1]
    else:
        for t in tables:
            if t.getAttribute("name") == sheet_spec:
                target_table = t
                break

    if target_table is None:
        sheet_names = [t.getAttribute("name") or f"Sheet{i}" for i, t in enumerate(tables, start=1)]
        print(f"Error: sheet '{sheet_spec}' not found (available: {', '.join(sheet_names)})", file=sys.stderr)
        sys.exit(1)

    for r in target_table.getElementsByType(table.TableRow):
        cells = [teletype.extractText(c).strip() for c in r.getElementsByType(table.TableCell)]
        all_rows.append(cells)
else:
    try:
        with open(wb_path, "rb") as handle:
            wb = openpyxl.load_workbook(handle, data_only=True, read_only=True)
            sheet_names = wb.sheetnames
            target_name = None

            if sheet_spec.isdigit():
                idx = int(sheet_spec)
                if 1 <= idx <= len(sheet_names):
                    target_name = sheet_names[idx - 1]
            elif sheet_spec in sheet_names:
                target_name = sheet_spec

            if target_name is None:
                print(f"Error: sheet '{sheet_spec}' not found (available: {', '.join(sheet_names)})", file=sys.stderr)
                sys.exit(1)

            ws = wb[target_name]
            for row in ws.iter_rows(values_only=True):
                all_rows.append(["" if c is None else str(c) for c in row])
    except Exception as e:
        print(f"Error: failed to open workbook: {e}", file=sys.stderr)
        sys.exit(1)

# Trim trailing completely empty rows
while all_rows and not any(str(c).strip() for c in all_rows[-1]):
    all_rows.pop()

if not all_rows:
    print("Warning: sheet is completely empty", file=sys.stderr)

out_dir = "/output" if os.path.exists("/output") else "."

if fmt in ("csv", "tsv"):
    delim = "\t" if fmt == "tsv" else ","
    out_buf = io.StringIO()
    writer = csv.writer(out_buf, delimiter=delim)
    writer.writerows(all_rows)
    content = out_buf.getvalue()

    file_name = f"sheet.{fmt}"
    out_path = os.path.join(out_dir, file_name)
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(content)

    sys.stdout.write(content)
elif fmt == "json":
    if len(all_rows) < 2:
        records = all_rows
    else:
        headers = [c.strip() or f"col_{i+1}" for i, c in enumerate(all_rows[0])]
        records = []
        for r in all_rows[1:]:
            padded = r + [""] * (len(headers) - len(r)) if len(r) < len(headers) else r[:len(headers)]
            records.append(dict(zip(headers, padded)))

    content = json.dumps(records, indent=2) + "\n"
    out_path = os.path.join(out_dir, "sheet.json")
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(content)

    sys.stdout.write(content)
